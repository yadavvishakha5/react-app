#!/usr/bin/env python3
"""Generate a professional DOCX resume for Vishakha Yadav."""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_cell_shading(cell, color_hex):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color_hex)
    shading.set(qn("w:val"), "clear")
    cell._tc.get_or_add_tcPr().append(shading)


def add_horizontal_line(paragraph):
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1F4E79")
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_section_heading(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(31, 78, 121)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    add_horizontal_line(p)
    return p


def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_after = Pt(2)
    if bold_prefix:
        run_b = p.add_run(bold_prefix)
        run_b.bold = True
        run_b.font.size = Pt(10)
        run = p.add_run(text)
        run.font.size = Pt(10)
    else:
        run = p.add_run(text)
        run.font.size = Pt(10)
    return p


doc = Document()

for section in doc.sections:
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

# Header
name = doc.add_paragraph()
name.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = name.add_run("Vishakha Yadav")
run.bold = True
run.font.size = Pt(20)
run.font.color.rgb = RGBColor(31, 78, 121)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("Senior Systems Engineer | React Native Developer")
run.font.size = Pt(11)
run.font.color.rgb = RGBColor(68, 68, 68)

contact = doc.add_paragraph()
contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = contact.add_run(
    "Pune, Maharashtra  |  +91 8408904267  |  vishakhayadav5000@gmail.com  |  LinkedIn  |  Portfolio"
)
run.font.size = Pt(9)
run.font.color.rgb = RGBColor(80, 80, 80)

# Summary
add_section_heading(doc, "Professional Summary")
summary = doc.add_paragraph(
    "React Native Developer with 4.5+ years of experience building secure, high-performance mobile banking "
    "applications for Android and iOS. Currently contributing to Union Bank of India mobile banking through IBM, "
    "with hands-on delivery across feature development, CR implementations, and accessibility enhancements. "
    "Strong expertise in React Native, TypeScript, Redux Toolkit, Finacle integration, and WCAG-compliant mobile UX. "
    "Proven track record in fintech delivery, change request ownership, and cross-functional collaboration in Agile environments."
)
summary.runs[0].font.size = Pt(10)

# Experience
add_section_heading(doc, "Professional Experience")

# Alchemy / IBM role
role = doc.add_paragraph()
r1 = role.add_run("Senior Systems Engineer (React Native Developer)")
r1.bold = True
r1.font.size = Pt(10.5)
role.add_run("\n")
r2 = role.add_run("Alchemy Techsol India Pvt. Ltd. (Payroll) — Deputed to IBM")
r2.bold = True
r2.font.size = Pt(10)
role.add_run("  |  March 2026 – Present  |  Pune, Maharashtra\n")
for run in role.runs[2:]:
    run.font.size = Pt(9.5)
    run.italic = True

client = doc.add_paragraph()
cr = client.add_run("Client / Project: ")
cr.bold = True
cr.font.size = Pt(10)
cr2 = client.add_run("Union Bank of India — Mobile Banking Applications")
cr2.bold = True
cr2.font.size = Pt(10)
cr2.font.color.rgb = RGBColor(31, 78, 121)

alchemy_bullets = [
    ("Develop and customize ", "secure, cross-platform mobile banking features for Android and iOS using React Native and TypeScript."),
    ("Delivered ", "2–3 Change Requests (CRs) end-to-end — from solution design and impact analysis through development, testing, and production release."),
    ("Built and customized the ", "Add/Edit Nominee flow, including UI implementation, validation logic, API integration, and alignment with banking business rules."),
    ("Drive ", "WCAG-compliant accessibility improvements across production mobile applications — enhancing TalkBack/VoiceOver support, navigation flows, semantic labeling, and inclusive UI patterns."),
    (None, "Collaborate with IBM architects, designers, QA, and business stakeholders to translate requirements into scalable, maintainable mobile solutions."),
    (None, "Integrate mobile features with backend banking services, ensuring secure data handling and seamless user experience across channels."),
]
for prefix, text in alchemy_bullets:
    add_bullet(doc, text, prefix)

# Infosys Senior
role2 = doc.add_paragraph()
role2.paragraph_format.space_before = Pt(6)
r1 = role2.add_run("Senior Systems Engineer")
r1.bold = True
r1.font.size = Pt(10.5)
role2.add_run("\n")
r2 = role2.add_run("Infosys Limited")
r2.bold = True
r2.font.size = Pt(10)
role2.add_run("  |  November 2023 – March 2026")
for run in role2.runs[2:]:
    run.font.size = Pt(9.5)
    run.italic = True

infosys_senior = [
    "Developed secure, scalable React Native applications using modular architecture, TypeScript, and Redux.",
    "Translated Figma designs into maintainable code using React Navigation and modern frontend practices.",
    "Implemented SSL pinning, secure storage, and device binding per enterprise banking standards.",
    "Integrated applications with Finacle Core Banking systems for real-time data sync and secure transactions.",
    "Designed RESTful APIs and FCM notification systems with token lifecycle management.",
    "Implemented Visa SDK in-app provisioning and Wallet Extensions for digital wallet integration.",
    "Improved CI/CD pipelines using GitHub Actions, CodePush, and Crashlytics.",
    "Prepared CR solution design documents with impact analysis and effort estimation.",
]
for b in infosys_senior:
    add_bullet(doc, b)

# Infosys Systems Engineer
role3 = doc.add_paragraph()
role3.paragraph_format.space_before = Pt(6)
r1 = role3.add_run("Systems Engineer")
r1.bold = True
r1.font.size = Pt(10.5)
role3.add_run("\n")
r2 = role3.add_run("Infosys Limited")
r2.bold = True
r2.font.size = Pt(10)
role3.add_run("  |  September 2021 – October 2023")
for run in role3.runs[2:]:
    run.font.size = Pt(9.5)
    run.italic = True

infosys_junior = [
    "Developed features for Finacle Mobile Banking (MB) 11.2.x platform.",
    "Built Cordova hybrid apps using IBM Worklight (MobileFirst) for Android and iOS.",
    "Integrated with Finacle Core Banking via secure APIs for real-time transaction processing.",
    "Implemented mobile security: encryption, tokenization, and ProGuard/DexGuard obfuscation.",
]
for b in infosys_junior:
    add_bullet(doc, b)

# Skills
add_section_heading(doc, "Technical Skills")

skills_data = [
    ("Mobile:", "React Native, Android, iOS, Hybrid Apps, Redux Toolkit, React Hooks"),
    ("Frontend:", "JavaScript (ES6+), TypeScript, HTML5, CSS3, AngularJS, React JS"),
    ("Banking:", "Finacle MB v11.x, Digital Banking, Nominee Management, Transaction Processing"),
    ("Security:", "SSL Pinning, Encryption, Device Binding, Secure Storage, ProGuard"),
    ("Accessibility:", "WCAG Compliance, TalkBack, VoiceOver, Inclusive UI/UX"),
    ("Tools:", "Git, JIRA, Postman, Android Studio, Xcode, VS Code, Firebase, Crashlytics"),
]
for label, items in skills_data:
    p = doc.add_paragraph()
    r1 = p.add_run(label + " ")
    r1.bold = True
    r1.font.size = Pt(10)
    r2 = p.add_run(items)
    r2.font.size = Pt(10)
    p.paragraph_format.space_after = Pt(1)

# Certifications
add_section_heading(doc, "Certifications & Achievements")
for item in [
    "ES6 JavaScript Certification",
    "Received the most prestigious award in the company",
    "Spot Award for exceptional project delivery and team collaboration",
]:
    add_bullet(doc, item)

# Education
add_section_heading(doc, "Education")
table = doc.add_table(rows=3, cols=3)
table.style = "Table Grid"
headers = ["Qualification", "Institution", "Score"]
for i, h in enumerate(headers):
    cell = table.rows[0].cells[i]
    cell.text = h
    set_cell_shading(cell, "1F4E79")
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)

edu_rows = [
    ("B.E. Electrical Engineering (2018–2021)", "AISSMS College of Engineering, Pune", "9.38 CGPA"),
    ("Diploma in Electrical Engineering (2015–2018)", "Government Polytechnic, Karad", "88.38%"),
]
for row_idx, row_data in enumerate(edu_rows, start=1):
    for col_idx, val in enumerate(row_data):
        cell = table.rows[row_idx].cells[col_idx]
        cell.text = val
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(9)

output_path = "/workspace/Vishakha_Yadav_Resume.docx"
doc.save(output_path)
print(f"Resume saved to {output_path}")
